import openpyxl
from openpyxl import Workbook
from io import BytesIO


def register_import_routes(app, get_db, render_template, request, redirect, url_for, flash, send_file):
    @app.route('/import_excel', methods=['GET', 'POST'])
    def import_excel():
        if request.method == 'POST':
            file = request.files.get('excel_file')
            if not file or not file.filename.endswith(('.xlsx', '.xls')):
                flash('الرجاء اختيار ملف Excel صالح', 'danger')
                return redirect(url_for('import_excel'))

            try:
                wb = openpyxl.load_workbook(file)
                ws = wb.active
                conn = get_db()
                success = 0
                failed = 0
                errors = []

                def fmt_date(d):
                    if d is None:
                        return None
                    if hasattr(d, 'strftime'):
                        return d.strftime('%Y-%m-%d')
                    return str(d).strip() or None

                for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row or not row[0]:
                        continue
                    try:
                        full_name = str(row[0]).strip()
                        age = row[1] if len(row) > 1 else None
                        gender = row[2] if len(row) > 2 else None
                        address = row[3] if len(row) > 3 else None
                        phone = row[4] if len(row) > 4 else None
                        disease_type = row[5] if len(row) > 5 else None
                        treatment_start = row[6] if len(row) > 6 else None
                        treatment_end = row[7] if len(row) > 7 else None
                        last_visit = row[8] if len(row) > 8 else None
                        status = row[9] if len(row) > 9 else 'قيد العلاج'
                        notes = row[10] if len(row) > 10 else None

                        conn.execute('''
                            INSERT INTO patients
                            (full_name, age, gender, address, phone, disease_type,
                             treatment_start, treatment_end, last_visit, status, notes)
                            VALUES (?,?,?,?,?,?,?,?,?,?,?)
                        ''', (
                            full_name, age, gender, address, phone, disease_type,
                            fmt_date(treatment_start), fmt_date(treatment_end),
                            fmt_date(last_visit), status or 'قيد العلاج', notes
                        ))
                        success += 1
                    except Exception as e:
                        failed += 1
                        errors.append(f"السطر {row_idx}: {str(e)}")

                conn.commit()
                conn.close()
                flash(f'تم الاستيراد بنجاح: {success} مريض. فشل: {failed}.',
                      'success' if failed == 0 else 'warning')
                for err in errors[:5]:
                    flash(err, 'danger')
                return redirect(url_for('admin'))
            except Exception as e:
                flash(f'خطأ في قراءة الملف: {str(e)}', 'danger')
                return redirect(url_for('import_excel'))

        return render_template('import_excel.html')

    @app.route('/download_template')
    def download_template():
        wb = Workbook()
        ws = wb.active
        ws.title = "المرضى"
        headers = ['الاسم الكامل', 'العمر', 'الجنس', 'العنوان', 'رقم الهاتف',
                   'نوع المرض', 'تاريخ بدء العلاج', 'تاريخ نهاية العلاج',
                   'تاريخ آخر زيارة', 'الحالة', 'ملاحظات']
        ws.append(headers)
        ws.append(['محمد أحمد علي', 35, 'ذكر', 'الرمادي - الملعب', '07701234567',
                   'رئوي', '2026-01-15', '', '', 'قيد العلاج', 'ملاحظة تجريبية'])
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True,
                         download_name='قالب_المرضى.xlsx')
